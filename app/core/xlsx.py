"""엑셀(.xlsx) 쓰기 — 도메인을 모르는 얇은 층 (D-403).

## 왜 openpyxl 인가

읽기는 표준 라이브러리로 한다(D-258 · `app/ai/parsing/ooxml.py`) — 그 결정의 근거는
「폐쇄망 wheelhouse 에 패키지를 늘리지 않는다」였다. 쓰기는 사용자가 외부 라이브러리를
허용하면서(2026-09-02) 검증된 구현을 쓰는 쪽으로 뒤집었다: openpyxl 과 전이 의존
et_xmlfile 은 둘 다 순수 파이썬 wheel 이라 wheelhouse 요구가 그대로 충족된다.
읽기 경로는 안 바꾼다 — 뒤집은 것은 쓰기 하나다.

## 수식 주입을 여기서 막는다

openpyxl 은 `=` 로 시작하는 문자열을 **수식으로 추론한다**(`cell.value = "=1+1"` 이
`data_type == "f"` 가 된다). 목록의 셀 값은 사용자가 적은 제목·이름이라, 그대로 두면
`=HYPERLINK(...)` 같은 제목이 받은 사람의 엑셀에서 **실행**된다. 모든 문자열 셀의
`data_type` 을 `"s"` 로 강제해 수식 추론을 끈다 — CSV 의 작은따옴표 접두(`csv_safe`)와
달리 **값 원문이 그대로 보존**된다(xlsx 에서 수식은 타입이 수식일 때만 수식이다).
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

__all__ = ["Sheet", "safe_sheet_title", "build_xlsx"]

#: 엑셀의 시트 이름 규칙 — 31자 상한, 일부 글자 금지. 어기면 파일이 열릴 때 복구 대화가 뜬다.
_TITLE_BAD = set("[]:*?/\\")
_TITLE_MAX = 31


@dataclass(frozen=True)
class Sheet:
    """시트 하나. `rows` 의 값은 str·int·float·None 이다 — 날짜는 부르는 쪽이 화면과
    같은 어휘(ISO 문자열)로 만들어 넘긴다(서식 스타일을 여기서 다시 발명하지 않는다)."""

    title: str
    headers: tuple[str, ...]
    rows: list[tuple] = field(default_factory=list)
    #: 열 폭(문자 수 단위). 비우면 openpyxl 기본 폭이다.
    widths: tuple[float, ...] = ()


def safe_sheet_title(title: str) -> str:
    """금지 글자를 걷고 31자에서 끊는다. 빈 이름은 「시트」다."""
    clean = "".join(ch for ch in (title or "") if ch not in _TITLE_BAD).strip()
    return (clean or "시트")[:_TITLE_MAX]


def _write_cell(ws, row_idx: int, col_idx: int, value: Any) -> None:
    cell = ws.cell(row=row_idx, column=col_idx)
    if value is None:
        return
    if isinstance(value, bool):
        # bool 은 int 의 하위형이라 아래 분기보다 먼저 잡는다 — 참/거짓을 숫자로 내보내면
        # 받은 쪽에서 1/0 이 무슨 뜻인지 알 수 없다.
        cell.value = "예" if value else "아니요"
        cell.data_type = "s"
        return
    if isinstance(value, (int, float)):
        cell.value = value
        return
    text = str(value)
    cell.value = text
    # 수식 추론을 끈다 — 파일 머리 주석의 그 문장이다.
    cell.data_type = "s"


def build_xlsx(sheets: Sequence[Sheet]) -> bytes:
    """시트 여럿을 한 통에 담아 bytes 로 돌려준다.

    머리글은 굵게, 첫 데이터 행에 틀 고정(freeze A2) — 긴 목록을 굴려도 열 이름이 남는다.
    """
    wb = Workbook()
    # openpyxl 은 기본 시트 하나를 들고 시작한다 — 첫 Sheet 가 그 자리를 쓴다.
    default = wb.active
    bold = Font(bold=True)

    for index, sheet in enumerate(sheets):
        ws = default if index == 0 else wb.create_sheet()
        ws.title = safe_sheet_title(sheet.title)
        for col_idx, header in enumerate(sheet.headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = str(header)
            cell.data_type = "s"
            cell.font = bold
        ws.freeze_panes = "A2"
        for row_offset, row in enumerate(sheet.rows, start=2):
            for col_idx, value in enumerate(row, start=1):
                _write_cell(ws, row_offset, col_idx, value)
        for col_idx, width in enumerate(sheet.widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = float(width)

    if not sheets:
        # 시트가 하나도 없으면 기본 시트에 이름만 준다 — 빈 통도 열리는 파일이어야 한다.
        default.title = safe_sheet_title("")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
