"""`GET /api/projects/{id}/export.xlsx` — 파일이 화면과 같은 사실을 말하는가 (D-403).

여기서 못박는 것:

1. 응답이 진짜 xlsx 이고(Content-Type·파일명·헤더), openpyxl 로 다시 열린다.
2. **WBS 시트가 화면의 `/wbs` 와 같은 표본**이다 — 계층 들여쓰기·상태·진행률.
   진행률 `None` 은 빈 칸이다(0 이 아니다).
3. `unplaced`(순환)도 담긴다 — 화면이 조용히 안 빼는 것과 같은 원칙.
4. 「기한 넘김」은 `milestone_unmet` 판정 그대로다 — 놓침+과거 = 넘김, 완료+과거 = 아님.
5. 감사 로그가 남는다(누가 언제 무엇을 내보냈나).
6. 상한을 넘으면 잘림을 **파일 안에서** 말하고 파일명에 -partial 이 붙는다.
"""

from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from app.core.models_base import join_names
from app.org.constants import DEFAULT_ORG_ID
from app.projects.models import Project, ProjectMilestone
from app.tickets.models import TicketCache
from app.work import relations

pytestmark = pytest.mark.integration

NOW = datetime(2026, 7, 14, 0, 0, 0)
BOSS_EMAIL = "prj-export@goodmit.co.kr"
PROJ_PAGE = "page-export"


def _ticket(db, *, page_id, title="작업", status="진행", est=None, parent=None):
    row = TicketCache(
        id=page_id, notion_page_id=page_id, org_id=DEFAULT_ORG_ID,
        title=title, status=status, est_wd=est, parent_page_id=parent,
        project_ids=join_names([PROJ_PAGE]), project_names="",
        assignee_notion_ids="", source="notion",
        synced_at=NOW, created_at=NOW, updated_at=NOW,
    )
    db.add(row)
    return row


@pytest.fixture()
def world(db, make_user):
    boss = make_user(BOSS_EMAIL, role="admin", display_name="관리자")
    boss.admin_scope = "global"
    db.commit()

    project = Project(
        name="내보내기 프로젝트", code="EXPRTA", org_id=DEFAULT_ORG_ID,
        notion_page_id=PROJ_PAGE, status="active",
        starts_on=date(2026, 7, 1), ends_on=date(2026, 9, 30),
    )
    db.add(project)
    db.flush()
    db.add_all([
        # 놓침 + 과거 기한 = 기한 넘김. 완료 + 과거 기한 = 아님(끝난 것만 지켜진 것이다).
        ProjectMilestone(project_id=project.id, name="1차 오픈", due_on=date(2026, 7, 10),
                         status="missed", sort_order=1),
        ProjectMilestone(project_id=project.id, name="설계 완료", due_on=date(2026, 7, 5),
                         status="done", sort_order=2),
        ProjectMilestone(project_id=project.id, name="2차 오픈", due_on=date(2026, 8, 20),
                         status="planned", sort_order=3),
    ])

    _ticket(db, page_id="ex-parent", title="상위 작업", est=100)
    _ticket(db, page_id="ex-child", title="자식 작업", est=4, status="완료", parent="ex-parent")
    _ticket(db, page_id="ex-solo", title="=HYPERLINK(\"https://evil\",\"x\")", est=6)
    db.flush()
    relations.sync_parent_links(db)
    db.commit()
    return {"project": project.id}


def _hdr(login_as):
    return {"X-CSRF-Token": login_as("admin", email=BOSS_EMAIL)}


def _export(client, login_as, project_id):
    r = client.get(f"/api/projects/{project_id}/export.xlsx", headers=_hdr(login_as))
    assert r.status_code == 200, f"{r.status_code} {r.text[:300]}"
    return r


def test_the_response_is_a_real_workbook_with_three_sheets(client, login_as, world):
    r = _export(client, login_as, world["project"])
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="project-EXPRTA-' in r.headers["content-disposition"]
    assert r.headers["cache-control"] == "no-store"
    assert r.headers["x-export-truncated"] == "0"

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["개요", "WBS", "마일스톤"]
    # 파일명은 ASCII 뿐이라 한글 이름은 개요 시트가 든다.
    overview = {row[0].value: row[1].value for row in wb["개요"].iter_rows(min_row=2)}
    assert overview["프로젝트 이름"] == "내보내기 프로젝트"
    assert overview["상태"] == "진행"
    assert overview["기간"] == "2026-07-01 ~ 2026-09-30"
    assert overview["기한 넘긴 마일스톤"] == 1
    assert overview["내보낸 날짜"] == "2026-07-14"


def test_the_wbs_sheet_mirrors_the_screen_sample(client, login_as, world):
    r = _export(client, login_as, world["project"])
    ws = load_workbook(io.BytesIO(r.content))["WBS"]
    # 픽스처 티켓에는 채번이 없어 「티켓 키」 열이 비어 있다 — 제목으로 찾는다.
    rows = {str(row[2].value or "").strip(): row for row in ws.iter_rows(min_row=2)}

    parent = rows["상위 작업"]
    child = rows["자식 작업"]
    assert parent[0].value == 1 and child[0].value == 2, "깊이 열이 계층을 안 말한다"
    assert child[2].value == "  자식 작업", "자식 제목이 들여쓰기로 계층을 보이지 않는다"
    assert child[3].value == "완료"
    # 상위 작업의 하위 트리 진행률: 4 / 4 = 100.0 (리프만 센다).
    assert parent[5].value == 100.0
    # 수식처럼 생긴 제목이 **글자 그대로** 남는다 — 수식이 아니다.
    evil = next(row for key, row in rows.items() if key.startswith("=HYPERLINK"))
    assert evil[2].value.startswith("=HYPERLINK")
    assert evil[2].data_type != "f"


def test_unplaced_cycles_are_included_not_dropped(client, login_as, world, db):
    _ticket(db, page_id="ex-cy-a", title="순환A", parent="ex-cy-b")
    _ticket(db, page_id="ex-cy-b", title="순환B", parent="ex-cy-a")
    db.flush()
    relations.sync_parent_links(db)
    db.commit()

    r = _export(client, login_as, world["project"])
    ws = load_workbook(io.BytesIO(r.content))["WBS"]
    text = "\n".join(str(c.value) for row in ws.iter_rows(min_row=2) for c in row if c.value)
    assert "순환A" in text and "순환B" in text, "못 그린 작업이 파일에서 조용히 사라졌다"


def test_the_milestone_sheet_uses_the_shared_overdue_rule(client, login_as, world):
    r = _export(client, login_as, world["project"])
    ws = load_workbook(io.BytesIO(r.content))["마일스톤"]
    rows = {row[0].value: row for row in ws.iter_rows(min_row=2)}
    assert rows["1차 오픈"][2].value == "놓침"
    assert rows["1차 오픈"][4].value == "예"
    assert rows["설계 완료"][4].value in ("", None), "완료 + 과거 기한은 넘김이 아니다"
    assert rows["2차 오픈"][4].value in ("", None), "미래 기한은 넘김이 아니다"


def test_the_export_is_audited(client, login_as, world, db):
    _export(client, login_as, world["project"])
    from sqlalchemy import select

    from app.audit.models import AuditLog

    db.expire_all()
    row = db.execute(
        select(AuditLog).where(
            AuditLog.action == "project.export_xlsx",
            AuditLog.object_id == world["project"],
        )
    ).scalars().first()
    assert row is not None, "내보내기가 감사 로그에 안 남았다"


def test_truncation_is_declared_in_the_file_and_the_filename(client, login_as, world, monkeypatch):
    from app.projects import export as export_mod

    monkeypatch.setattr(export_mod, "EXPORT_MAX_WBS_ROWS", 1)
    r = _export(client, login_as, world["project"])
    assert r.headers["x-export-truncated"] == "1"
    assert "-partial-" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content))["WBS"]
    last = list(ws.iter_rows(min_row=2))[-1]
    assert "잘렸습니다" in str(last[2].value), "잘림을 파일이 스스로 말해야 한다"


def test_a_project_without_ascii_code_falls_back_to_id(client, login_as, db, make_user):
    project = Project(name="코드 없는 프로젝트", org_id=DEFAULT_ORG_ID, status="planned")
    db.add(project)
    db.commit()
    r = _export(client, login_as, project.id)
    assert f'filename="project-{project.id[:8]}-' in r.headers["content-disposition"]


def test_progress_none_is_an_empty_cell_not_zero(client, login_as, db, make_user):
    """작업이 하나도 없는 프로젝트 — 진행률은 0 이 아니라 빈 칸이다."""
    project = Project(name="빈 프로젝트", code="EMPTYA", org_id=DEFAULT_ORG_ID, status="planned")
    db.add(project)
    db.commit()
    r = _export(client, login_as, project.id)
    wb = load_workbook(io.BytesIO(r.content))
    overview = {row[0].value: row[1].value for row in wb["개요"].iter_rows(min_row=2)}
    assert overview["진행률(%)"] is None
    assert overview["진행률 근거"] == "집계할 작업이 없습니다"
