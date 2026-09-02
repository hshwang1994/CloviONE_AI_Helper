"""`app/core/xlsx.py` — 만든 파일이 실제로 열리고, 수식 주입이 막히는가 (D-403).

## Known Good 을 두 갈래로 잰다

1. **openpyxl 왕복** — 쓴 것을 같은 라이브러리로 다시 연다. 시트 이름·순서·값·타입.
2. **제품 리더 교차검증** — `app/ai/parsing/ooxml.py::parse_xlsx`(표준 라이브러리, rels 를
   따라가는 독립 구현)가 같은 파일을 읽는다. 쓰기와 읽기가 서로 다른 구현이라, 이것이
   통과하면 「우리 라이브러리끼리만 통하는 파일」이 아니라는 뜻이다.

## 수식 주입

openpyxl 은 `=` 시작 문자열을 수식으로 추론한다 — 사용자 제목이 `=HYPERLINK(...)` 면
받은 사람의 엑셀에서 실행된다. 문자열 셀의 `data_type` 을 `"s"` 로 강제해 그 추론을
껐고, CSV 의 작은따옴표 접두와 달리 **값 원문이 보존**되는지도 함께 본다.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.core.xlsx import Sheet, build_xlsx, safe_sheet_title

pytestmark = pytest.mark.unit


def _load(body: bytes):
    return load_workbook(io.BytesIO(body))


def test_a_roundtrip_keeps_sheets_headers_and_values():
    body = build_xlsx([
        Sheet("개요", ("항목", "값"), [("이름", "알파"), ("점수", 87)]),
        Sheet("WBS", ("깊이", "제목"), [(1, "상위"), (2, "  하위")]),
    ])
    wb = _load(body)
    assert wb.sheetnames == ["개요", "WBS"]
    ws = wb["개요"]
    assert [c.value for c in ws[1]] == ["항목", "값"]
    assert ws.cell(row=2, column=1).value == "이름"
    assert ws.cell(row=3, column=2).value == 87
    # 들여쓰기 공백이 살아 있다 — 트리를 편 제목의 계층 표기다.
    assert wb["WBS"].cell(row=3, column=2).value == "  하위"


def test_numbers_stay_numbers_and_none_stays_empty():
    body = build_xlsx([Sheet("표", ("이름", "값"), [("정수", 3), ("실수", 6.5), ("없음", None)])])
    ws = _load(body)["표"]
    assert ws.cell(row=2, column=2).value == 3
    assert ws.cell(row=3, column=2).value == 6.5
    # None 은 빈 칸이다 — 0 으로 바꾸면 「없음」과 「0」이 같아 보인다(null≠0 계약).
    assert ws.cell(row=4, column=2).value is None


@pytest.mark.parametrize("payload", [
    "=HYPERLINK(\"https://evil.example\",\"클릭\")",
    "=1+1",
    "+1+2",
    "-2+3",
    "@SUM(A1)",
    "\t=cmd",
])
def test_formula_looking_strings_stay_literal_text(payload):
    """🔴 수식으로 실행되지 않고, 값 원문이 그대로 남는다."""
    body = build_xlsx([Sheet("표", ("제목",), [(payload,)])])
    ws = _load(body)["표"]
    cell = ws.cell(row=2, column=1)
    assert cell.data_type != "f", "수식으로 저장됐다 — 받은 사람의 엑셀에서 실행된다"
    assert cell.value == payload, "방어가 값을 훼손했다 — 원문이 보존돼야 한다"


def test_special_characters_survive():
    tricky = "한글 <태그> & \"따옴표\" '작은' 줄\n바꿈"
    body = build_xlsx([Sheet("표", ("값",), [(tricky,)])])
    assert _load(body)["표"].cell(row=2, column=1).value == tricky


def test_sheet_titles_are_sanitised():
    assert safe_sheet_title("a[b]c:d*e?f/g\\h") == "abcdefgh"
    assert len(safe_sheet_title("가" * 40)) == 31
    assert safe_sheet_title("   ") == "시트"
    # 통 안에서도 안전한 이름으로 저장된다.
    body = build_xlsx([Sheet("나쁜:이름?", ("값",), [("x",)])])
    assert _load(body).sheetnames == ["나쁜이름"]


def test_an_empty_workbook_still_opens():
    wb = _load(build_xlsx([]))
    assert len(wb.sheetnames) == 1


def test_booleans_become_korean_words():
    body = build_xlsx([Sheet("표", ("켜짐",), [(True,), (False,)])])
    ws = _load(body)["표"]
    assert ws.cell(row=2, column=1).value == "예"
    assert ws.cell(row=3, column=1).value == "아니요"


def test_the_zip_has_the_required_ooxml_parts():
    """xlsx 는 zip 이다 — 필수 파트가 없으면 엑셀이 복구 대화상자를 띄운다."""
    body = build_xlsx([Sheet("표", ("값",), [("x",)])])
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        names = set(archive.namelist())
    assert "[Content_Types].xml" in names
    assert "xl/workbook.xml" in names
    assert "xl/_rels/workbook.xml.rels" in names
    assert any(n.startswith("xl/worksheets/sheet") for n in names)


def test_the_product_reader_can_read_what_we_wrote(tmp_path: Path):
    """교차검증 — 쓰기(openpyxl)와 **다른 구현**(stdlib 리더)이 같은 내용을 읽는다."""
    from app.ai.parsing.ooxml import parse_xlsx

    body = build_xlsx([
        Sheet("회의록", ("제목", "상태"), [("킥오프 정리", "완료"), ("설계 검토", "진행")]),
        Sheet("마일스톤", ("이름",), [("1차 오픈",)]),
    ])
    path = tmp_path / "roundtrip.xlsx"
    path.write_bytes(body)

    result = parse_xlsx(path)
    assert result.ok, getattr(result, "detail", result)
    joined = "\n".join(unit.text for unit in result.units)
    assert "킥오프 정리" in joined
    assert "1차 오픈" in joined
    labels = [unit.anchor_label for unit in result.units]
    assert labels == ["회의록", "마일스톤"]
